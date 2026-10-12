"""
Export everything to one Excel workbook: logs/tollgaterag_metrics.xlsx

Sheets: Provenance, Query Summary, Risk Distribution, Policy Comparison,
        Ablation, Threshold Sweep, Per-Document, Review Queue.
(Run `python -m tollgaterag.eval` and `python -m tollgaterag.ablation` first.)
"""

import json
import os
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from tollgaterag import config
from tollgaterag.provenance import load_log, load_review_queue

OUTPUT_PATH = os.path.join(config.LOG_DIR, "tollgaterag_metrics.xlsx")
FILL = {"accept": PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE"),
        "warn": PatternFill("solid", start_color="FFEB9C", end_color="FFEB9C"),
        "reject": PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")}


def _autosize(ws):
    for col in ws.columns:
        w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, 60)


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        cleaned_row = [str(v) if isinstance(v, (list, dict)) else v for v in r]
        ws.append(cleaned_row)
    _autosize(ws)
    return ws


def _load_json(name):
    p = os.path.join(config.LOG_DIR, name)
    return json.load(open(p)) if os.path.exists(p) else None


def build_workbook(output_path=OUTPUT_PATH):
    records = load_log()
    if not records:
        raise RuntimeError("No provenance log yet. Run `python -m tollgaterag.eval` first.")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Provenance
    rows = []
    for rec in records:
        used = set(rec["sources_used_in_answer"])
        for c in rec["candidates_evaluated"]:
            s = c["sub_scores"]
            rows.append([rec["query_id"], rec["timestamp"], rec["query"], c["doc_id"], c["title"],
                         c["overall_risk"], s["trust_level"], s["freshness"], s["injection"],
                         s["authority_claim"], s["anchor_conflict"], c["decision"], c["doc_id"] in used])
    ws = _sheet(wb, "Provenance",
                ["query_id", "timestamp", "query", "doc_id", "title", "overall_risk", "trust", "freshness",
                 "injection", "authority_claim", "anchor_conflict", "decision", "used_in_answer"], rows)
    for r in range(2, ws.max_row + 1):
        f = FILL.get(ws.cell(row=r, column=12).value)
        if f:
            ws.cell(row=r, column=12).fill = f

    # Query summary
    rows = []
    for rec in records:
        cnt = defaultdict(int)
        for c in rec["candidates_evaluated"]:
            cnt[c["decision"]] += 1
        rows.append([rec["query_id"], rec["query"], rec["latency_ms"], cnt["accept"], cnt["warn"],
                     cnt["reject"], rec["answer"][:120].replace("\n", " ")])
    _sheet(wb, "Query Summary", ["query_id", "query", "latency_ms", "accept", "warn", "reject", "answer_preview"], rows)

    # Risk distribution
    allc = [c for rec in records for c in rec["candidates_evaluated"]]
    n = len(allc)
    cnt, sub = defaultdict(int), defaultdict(float)
    for c in allc:
        cnt[c["decision"]] += 1
        for k, v in c["sub_scores"].items():
            sub[k] += v
    rows = [["Total candidate evaluations", n], ["Accepted", cnt["accept"]], ["Warned", cnt["warn"]],
            ["Rejected", cnt["reject"]]] + [[f"avg_{k}_risk", round(v / n, 4)] for k, v in sub.items()]
    _sheet(wb, "Risk Distribution", ["Metric", "Value"], rows)

    ev = _load_json("eval_test_split.json") or _load_json("eval_results.json")
    if ev and "policies" in ev:
        keys = ["adv_reject_rate", "adv_flag_rate", "benign_reject_rate", "benign_flag_rate",
                "low_quality_flag_rate", "verified_reject_rate", "contamination",
                "contamination_plain_rag", "adv_docs_in_context", "adv_docs_in_context_plain_rag"]
        _sheet(wb, "Policy Comparison", ["policy"] + keys,
               [[n] + [m.get(k, 0.0) for k in keys] for n, m in ev["policies"].items()])
        if "per_doc" in ev:
            _sheet(wb, "Per-Document", ["doc_id", "label", "events", "rejected", "flagged"],
                   [[d, v["label"], v["events"], v["rejected"], v["flagged"]] for d, v in sorted(ev["per_doc"].items())])

    ab = _load_json("ablation_results.json")
    if ab:
        if "leave_one_out_ablation" in ab:
            rows = ab["leave_one_out_ablation"]
            hdr = [h for h in rows[0].keys()]
            _sheet(wb, "Leave-One-Out Ablation", hdr, [[r.get(h) for h in hdr] for r in rows])
        elif "ablation" in ab:
            rows = ab["ablation"]
            hdr = [h for h in rows[0].keys()]
            _sheet(wb, "Ablation", hdr, [[r.get(h) for h in hdr] for r in rows])

        if "architectural_ablations" in ab:
            rows = ab["architectural_ablations"]
            hdr = [h for h in rows[0].keys()]
            _sheet(wb, "Architectural Ablation", hdr, [[r.get(h) for h in hdr] for r in rows])

        if "validation_threshold_sweep" in ab:
            rows = ab["validation_threshold_sweep"]
            hdr = [h for h in rows[0].keys()]
            _sheet(wb, "Validation Threshold Sweep", hdr, [[r.get(h) for h in hdr] for r in rows])
        elif "threshold_sweep" in ab:
            rows = ab["threshold_sweep"]
            hdr = [h for h in rows[0].keys()]
            _sheet(wb, "Threshold Sweep", hdr, [[r.get(h) for h in hdr] for r in rows])

    rq = load_review_queue()
    if rq:
        _sheet(wb, "Review Queue", ["timestamp", "query", "doc_id", "title", "decision", "overall_risk", "status"],
               [[r.get("timestamp"), r.get("query"), r.get("doc_id"), r.get("title"), r.get("decision"), r.get("overall_risk"), r.get("status")]
                for r in rq])

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"Saved metrics workbook: {output_path}  (sheets: {', '.join(wb.sheetnames)})")
    return output_path


if __name__ == "__main__":
    build_workbook()
